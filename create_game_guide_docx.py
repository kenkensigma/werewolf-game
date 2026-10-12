import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, accent_color_hex="C6A66B", bg_color_hex="F7F5F0"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{accent_color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"PERINGATAN ATURAN: {title}\n" if title else "")
    run_title.font.name = "Georgia"
    run_title.font.size = Pt(10)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(141, 52, 59) if accent_color_hex == "8D343B" else RGBColor(198, 166, 107)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(40, 40, 40)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_document():
    doc = docx.Document()
    
    # Page setup: Standard A4
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Header and Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Werewolf Multiplayer : Gothic Social Deduction | Panduan Bermain Resmi")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(130, 130, 130)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_sub = p_title.add_run("PANDUAN RESMI BERMAIN\n")
    r_sub.font.name = "Georgia"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(198, 166, 107) # Gold
    
    r_main = p_title.add_run("WEREWOLF MULTIPLAYER")
    r_main.font.name = "Georgia"
    r_main.font.size = Pt(26)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(25, 26, 31)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_desc = p_sub.add_run("Gothic Social Deduction Game : Petunjuk Lengkap, Mekanika Peran, dan Strategi")
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(11)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(100, 100, 100)
    
    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r_div.font.size = Pt(8)
    r_div.font.color.rgb = RGBColor(198, 166, 107)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(141, 52, 59) # Burgundy
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 40, 60)
        return h

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(25, 25, 25)
        rt = p.add_run(text)
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.italic = italic
        rt.font.color.rgb = RGBColor(40, 40, 40)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(25, 25, 25)
        rt = p.add_run(text)
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = RGBColor(40, 40, 40)
        return p

    # SECTION 1
    add_h1("1. TENTANG PERMAINAN")
    add_body(
        "Werewolf Multiplayer adalah permainan sosial deduksi berbasis browser dengan tema Gothic Tabletop yang kelam dan imersif. "
        "Di dalam sebuah desa terpencil yang diselimuti kabut mistis, sebagian warga diam-diam terjangkit kutukan mengerikan dan berubah menjadi Manusia Serigala (Werewolf) di malam hari."
    )
    add_body(
        "Setiap pemain akan diberikan peran rahasia oleh sistem server secara acak dan terenkripsi. Pemain harus menggunakan kecerdasan, "
        "kemampuan analisa, diplomasi, serta deduksi logis untuk membawa faksi mereka menuju kemenangan."
    )
    add_body(
        "Tiga faksi utama yang bertarung di dalam permainan ini:", bold_prefix="Tiga Faksi Utama: "
    )
    add_bullet("Terdiri dari Warga Desa, Peramal (Seer), Penjaga (Guard), dan Pemburu (Hunter). Tujuan mereka adalah menemukan dan melenyapkan seluruh Manusia Serigala melalui musyawarah siang hari.", bold_prefix="Faksi Warga (Village): ")
    add_bullet("Terdiri dari kawanan Manusia Serigala yang saling berkoordinasi memangsa warga di malam hari. Tujuan mereka adalah mencapai jumlah paritas setara atau lebih banyak dari warga yang masih hidup.", bold_prefix="Faksi Serigala (Werewolf): ")
    add_bullet("Peran Pembuat Onar (Troublemaker). Faksi mandiri yang memiliki agenda tersembunyi untuk membuat kekacauan dan menang secara solo dengan memanipulasi pengadilan.", bold_prefix="Faksi Independen (Troublemaker): ")

    # SECTION 2
    add_h1("2. CARA MEMULAI DAN PENGATURAN KAMAR (LOBBY)")
    add_h2("2.1 Membuat Kamar Baru (Create Chamber)")
    add_bullet("Buka aplikasi permainan di browser.", bold_prefix="Langkah 1: ")
    add_bullet("Pilih tab 'Establish Chamber', masukkan Nama Tampilan (Display Name) Anda, lalu klik tombol 'Establish Chamber'.", bold_prefix="Langkah 2: ")
    add_bullet("Sistem server akan menghasilkan Kode Kamar unik sepanjang 6 karakter (contoh: UUZQKZ).", bold_prefix="Langkah 3: ")
    add_bullet("Klik tombol 'Copy Code' di bagian atas layar untuk membagikan kode kamar kepada teman-teman Anda.", bold_prefix="Langkah 4: ")

    add_h2("2.2 Bergabung ke Kamar yang Sudah Ada (Join Chamber)")
    add_bullet("Pilih tab 'Enter Chamber' pada halaman awal.", bold_prefix="Langkah 1: ")
    add_bullet("Ketikkan 6 karakter Kode Kamar yang diberikan oleh Tuan Rumah (Host).", bold_prefix="Langkah 2: ")
    add_bullet("Masukkan Nama Tampilan Anda dan tekan tombol 'Enter Chamber'.", bold_prefix="Langkah 3: ")

    add_h2("2.3 Pengaturan Kamar oleh Tuan Rumah (Host Settings)")
    add_body("Tuan Rumah (Host) memiliki wewenang untuk mengatur parameter ruangan sebelum memulai perburuan:")
    add_bullet("Menentukan jumlah serigala dalam permainan (default: 2 serigala).", bold_prefix="Jumlah Serigala (Werewolf Count): ")
    add_bullet("Mengaktifkan peran Pembuat Onar. Jika diaktifkan, batas minimal pemain akan otomatis bertambah menjadi 9 orang.", bold_prefix="Peran Troublemaker: ")
    add_bullet("Durasi waktu berdebat di siang hari (Pilihan: 2 Menit, 4 Menit standar, atau 6 Menit).", bold_prefix="Durasi Diskusi Siang: ")

    add_callout_box(
        doc,
        "SYARAT MINIMAL PEMAIN",
        "Mode Standar membutuhkan minimal 8 pemain (1 Seer, 2 Guard, 2 Hunter, 2 Werewolf, 1 Villager). "
        "Jika Troublemaker diaktifkan, dibutuhkan minimal 9 pemain. "
        "Untuk keperluan pengujian atau simulasi, Host dapat menekan tombol 'Summon Test Villagers (Dev)' untuk mengisi slot yang kosong secara instan.",
        accent_color_hex="C6A66B"
    )

    # SECTION 3
    add_h1("3. DAFTAR PERAN LENGKAP & KEKUATAN KHUSUS")
    
    # Table of Roles
    table_roles = doc.add_table(rows=7, cols=4)
    table_roles.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_roles.autofit = False
    
    col_widths = [Inches(1.3), Inches(1.0), Inches(2.5), Inches(1.7)]
    headers = ["Peran", "Faksi", "Kemampuan Khusus", "Kondisi Menang"]
    
    # Header Row
    hdr_cells = table_roles.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "191A1F")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Georgia"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(198, 166, 107)

    roles_data = [
        ("Werewolf\n(Serigala)", "Werewolf", "Berkoordinasi dengan kawanan setiap malam untuk memilih satu korban mangsa.", "Jumlah serigala hidup >= jumlah warga non-serigala hidup (Paritas)."),
        ("Seer\n(Peramal)", "Warga", "Meneropong 1 pemain per malam. Menerima kartu ramalan 3D privat dengan hasil: WEREWOLF atau NOT WEREWOLF.", "Seluruh Werewolf berhasil dilenyapkan dari desa."),
        ("Guard\n(Penjaga)", "Warga", "Melindungi 1 pemain hidup per malam dari serangan serigala dan panah pemburu. Tidak bisa melindungi diri sendiri.", "Seluruh Werewolf berhasil dilenyapkan dari desa."),
        ("Hunter\n(Pemburu)", "Warga", "Membawa panah perak di malam hari. Dapat membidik 1 pemain atau menahan tembakan (Hold Fire / Skip).", "Seluruh Werewolf berhasil dilenyapkan dari desa."),
        ("Villager\n(Warga)", "Warga", "Tidak memiliki kekuatan malam. Mengandalkan logika, debat, dan pemungutan suara di siang hari.", "Seluruh Werewolf berhasil dilenyapkan dari desa."),
        ("Troublemaker\n(Pembuat Onar)", "Independen", "Tidur di malam hari tanpa kemampuan aktif. Bertujuan membuat huru-hara di siang hari.", "HANYA MENANG jika dieksekusi di pengadilan siang hari!"),
    ]

    for row_idx, r_data in enumerate(roles_data, start=1):
        row_cells = table_roles.rows[row_idx].cells
        bg_col = "FFFFFF" if row_idx % 2 == 1 else "F9F8F6"
        for col_idx in range(4):
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(r_data[col_idx])
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if r_data[col_idx] == "Werewolf":
                    r.font.color.rgb = RGBColor(141, 52, 59)
                elif r_data[col_idx] == "Warga":
                    r.font.color.rgb = RGBColor(60, 90, 130)
                else:
                    r.font.color.rgb = RGBColor(180, 140, 40)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)

    # Detailed role explanations
    add_h2("3.1 Mekanika Khusus Peramal (Seer) & Kartu Ramalan 3D")
    add_body(
        "Sesuai standar kerahasiaan ketat pada dokumen desain game, Peramal tidak pernah menerima nama peran persis dari target. "
        "Hasil penyelidikan hanya dikirim secara privat ke layar Peramal dan bernilai tepat salah satu dari dua kategori:"
    )
    add_bullet("Jika target yang diselidiki memiliki peran Werewolf.", bold_prefix="WEREWOLF : ")
    add_bullet("Jika target yang diselidiki adalah Warga, Penjaga, Pemburu, Troublemaker, atau sesama Peramal.", bold_prefix="NOT WEREWOLF : ")
    add_body(
        "Peramal akan melihat animasi kartu 3D berputar (Mirror of Truth). Kartu berwana merah marun gelap menandakan Serigala, "
        "sedangkan kartu bercahaya biru rembulan menandakan jiwa yang bukan serigala.", italic=True
    )

    add_h2("3.2 Mekanika Khusus Pembuat Onar (Troublemaker)")
    add_body(
        "Troublemaker adalah peran independen yang eksentrik. Target utamanya adalah membuat dirinya dicurigai dan dihukum mati oleh warga di pengadilan desa."
    )
    add_bullet("Jika Troublemaker dieksekusi melalui pemungutan suara siang hari, permainan BERAKHIR SEKETIKA dan Troublemaker dinobatkan sebagai SATU-SATUNYA PEMENANG!", bold_prefix="Kemenangan Siang: ")
    add_bullet("Jika Troublemaker dibunuh di malam hari oleh Serigala atau Pemburu, Troublemaker TIDAK MENANG dan langsung tersingkir dari permainan.", bold_prefix="Kematian Malam: ")

    # SECTION 4
    add_h1("4. ALUR & FASE PERMAINAN LENGKAP")
    add_body(
        "Permainan berjalan melalui siklus berulang yang dikendalikan penuh oleh otoritas server secara aman:"
    )

    # Phases table
    table_phases = doc.add_table(rows=6, cols=3)
    table_phases.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_phases.autofit = False
    
    col_w_phases = [Inches(1.8), Inches(1.2), Inches(3.5)]
    hdr_phases = ["Fase Permainan", "Durasi Server", "Aktivitas & Aturan Utama"]
    
    hdr_p_cells = table_phases.rows[0].cells
    for i, title in enumerate(hdr_phases):
        hdr_p_cells[i].width = col_w_phases[i]
        set_cell_background(hdr_p_cells[i], "191A1F")
        set_cell_margins(hdr_p_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_p_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Georgia"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(198, 166, 107)

    phases_data = [
        ("0. Pengungkapan Peran\n(Role Reveal)", "12 Detik / Klik Siap", "Setiap pemain membalik kartu 3D privat untuk membaca peran dan kekuatan rahasianya. Klik 'Acknowledge Secret' jika sudah siap."),
        ("1. Malam Pertama\n(Quiet Shadows)", "6 Detik", "Malam hening pembuka. Tidak ada cakar serigala dan tidak ada aksi mematikan. Desa bersiap menyongsong fajar pertama."),
        ("2. Diskusi Siang\n(Day Discussion)", "2 - 6 Menit (Config)", "Seluruh warga yang masih hidup berdebat melalui obrolan (Deliberation Scroll). Warga yang sudah mati dilarang mengirim pesan."),
        ("3. Pengadilan Desa\n(Tribunal Voting)", "45 Detik / Selesai", "Pemungutan suara terbuka. Setiap warga dapat memilih 1 tersangka atau Skip. Calon dengan suara tertinggi mutlak akan dieksekusi."),
        ("4. Malam Perburuan\n(Dead of Night)", "30 Detik / Selesai", "Kekuatan malam aktif serentak. Serigala memilih mangsa, Peramal meneropong, Penjaga memasang perisai, Pemburu membidik."),
    ]

    for row_idx, p_data in enumerate(phases_data, start=1):
        row_cells = table_phases.rows[row_idx].cells
        bg_col = "FFFFFF" if row_idx % 2 == 1 else "F9F8F6"
        for col_idx in range(3):
            row_cells[col_idx].width = col_w_phases[col_idx]
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(p_data[col_idx])
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_spacer2 = doc.add_paragraph()
    p_spacer2.paragraph_format.space_before = Pt(8)

    add_h2("4.1 Aturan Ketat Pengadilan Siang Hari (Voting Resolution)")
    add_body(
        "Di dalam pengadilan desa, eliminasi ditentukan dengan prinsip Unique Highest Vote Count:"
    )
    add_bullet("Seorang tersangka hanya akan dieksekusi jika mendapatkan jumlah suara terbanyak secara mutlak (tidak seri dengan calon lain).", bold_prefix="Syarat Eksekusi: ")
    add_bullet("Jika dua calon teratas mendapatkan suara yang sama persis (misal: Calon A mendapat 3 suara, Calon B mendapat 3 suara), maka pengadilan DEADLOCK dan TIDAK ADA PEMAIN YANG DIEKSEKUSI.", bold_prefix="Hasil Seri (Tie): ")
    add_bullet("Pemain dapat memilih 'Abstain / Skip Vote' jika merasa belum ada bukti yang cukup kuat untuk menggantung seorang tersangka.", bold_prefix="Opsi Skip: ")
    add_bullet("Pemain tidak dapat memilih dirinya sendiri, dan tidak dapat memilih pemain yang sudah mati.", bold_prefix="Validasi Suara: ")

    add_h2("4.2 Mesin Resolusi Malam Simultan (Simultaneous Night Resolution)")
    add_body(
        "Salah satu keunggulan arsitektur game ini adalah sistem resolusi simultan pada fajar hari. "
        "Server mengumpulkan seluruh aksi malam terlebih dahulu sebelum menghitung eliminasi secara bersamaan:"
    )
    add_bullet("Server mengambil potret pemain yang hidup sebelum aksi malam diselesaikan.", bold_prefix="1. Snapshot Pemain Hidup : ")
    add_bullet("Jika beberapa serigala memilih target yang berbeda, target dengan suara terbanyak di dalam pack akan diserang. Jika seri di dalam pack, serangan serigala batal (skip).", bold_prefix="2. Konsensus Serigala : ")
    add_bullet("Pemain yang dilindungi oleh Penjaga (Guard) kebal terhadap seluruh serangan malam, termasuk cakar Werewolf dan panah Pemburu.", bold_prefix="3. Perisai Penjaga : ")
    add_bullet("Jika satu target diserang oleh Serigala sekaligus ditembak oleh Pemburu tanpa perlindungan, ia hanya dieliminasi tepat 1 kali tanpa duplikasi status.", bold_prefix="4. Deduplikasi Serangan : ")
    add_bullet("Jika tidak ada korban jiwa, banner 'Dawn Breaks in Peace' akan menyambut pagi hari. Jika ada korban, banner 'Casualty of the Night' akan mengumumkan nama pemain yang gugur.", bold_prefix="5. Pengumuman Fajar : ")

    # SECTION 5
    add_h1("5. STRATEGI DAN TIPS UNTUK SETIAP PERAN")
    
    add_h2("5.1 Tips untuk Warga Desa (Villager)")
    add_bullet("Perhatikan pola bicara dan pemilihan suara pemain lain. Serigala biasanya ragu-ragu atau mencoba mengalihkan tuduhan ke warga yang vokal.", bold_prefix="Amati Inkonsistensi: ")
    add_bullet("Jangan ragu memilih 'Skip' jika bukti belum jelas daripada mengeksekusi sesama warga secara terburu-buru.", bold_prefix="Manfaatkan Skip: ")

    add_h2("5.2 Tips untuk Manusia Serigala (Werewolf)")
    add_bullet("Jangan terlalu diam dan jangan terlalu agresif. Bersikaplah seperti warga biasa yang sedang mencari serigala.", bold_prefix="Membaur (Blend In): ")
    add_bullet("Singkirkan Peramal dan Penjaga sesegera mungkin di malam hari untuk melemahkan pertahanan desa.", bold_prefix="Target Prioritas: ")
    add_bullet("Jika rekan serigala Anda sudah terlanjur dicurigai oleh mayoritas warga, ikutlah memilihnya agar Anda tidak ikut dicurigai.", bold_prefix="Taktik Pengorbanan: ")

    add_h2("5.3 Tips untuk Peramal (Seer)")
    add_bullet("Jangan langsung mengumumkan diri Anda sebagai Peramal di hari pertama kecuali Anda berhasil menemukan Werewolf. Jika Anda membuka diri terlalu dini, serigala akan langsung membunuh Anda di malam berikutnya.", bold_prefix="Jaga Kerahasiaan: ")
    add_bullet("Jika Anda menemukan Werewolf, sampaikan informasi dengan hati-hati atau bangun aliansi rahasia dengan warga yang Anda percaya.", bold_prefix="Sampaikan Visi: ")

    add_h2("5.4 Tips untuk Penjaga (Guard)")
    add_bullet("Lindungi pemain yang paling berkontribusi atau pemain yang secara terbuka mengaku sebagai Peramal yang sah.", bold_prefix="Pemain Kunci: ")
    add_bullet("Karena Anda tidak bisa melindungi diri sendiri, usahakan agar gaya bermain Anda tidak terlalu menonjol sehingga serigala tidak mengincar Anda lebih dulu.", bold_prefix="Sembunyikan Identitas: ")

    add_h2("5.5 Tips untuk Pemburu (Hunter)")
    add_bullet("Gunakan opsi 'Hold Fire (Skip)' jika Anda belum yakin 100%. Menembak sesama warga hanya akan mempercepat kemenangan serigala.", bold_prefix="Disiplin Menembak: ")
    add_bullet("Simpan panah perak Anda untuk momen genting saat jumlah serigala dan warga mulai seimbang.", bold_prefix="Waktu Eksekusi: ")

    add_h2("5.6 Tips untuk Pembuat Onar (Troublemaker)")
    add_bullet("Buatlah argumen yang sedikit ganjil, sering berpindah pilihan, atau berpura-pura menjadi serigala yang ceroboh.", bold_prefix="Ciptakan Kecurigaan: ")
    add_bullet("Jangan mengaku secara terang-terangan bahwa Anda adalah Troublemaker, karena warga yang cerdas akan menolak mengeksekusi Anda di pengadilan!", bold_prefix="Jangan Terlalu Terbuka: ")

    # SECTION 6
    add_h1("6. FITUR AKHIR PERMAINAN & TANDING ULANG (REMATCH)")
    add_body(
        "Ketika kondisi kemenangan salah satu faksi terpenuhi, permainan memasuki fase Judgement Concluded (Game Over) dengan tampilan layar megah:"
    )
    add_bullet("Menampilkan faksi pemenang disertai lambang lambang heraldik resmi (Matahari Fajar untuk Warga, Serigala Malam untuk Werewolf, Mahkota Duri untuk Troublemaker).", bold_prefix="Pengumuman Kemenangan: ")
    add_bullet("Tirai kerahasiaan dibuka penuh! Seluruh identitas peran asli pemain diumumkan ke semua orang bersama status akhir (Selamat, Dieksekusi di Pengadilan, atau Gugur dalam Bayangan Malam).", bold_prefix="Tabel Pengungkapan Peran (The Veil Is Lifted): ")
    add_bullet("Rekapitulasi total siklus siang/malam, jumlah jiwa yang selamat, dan korban jiwa.", bold_prefix="Statistik Pertandingan: ")
    add_bullet("Tuan Rumah (Host) dapat menekan tombol 'Commence Rematch' untuk langsung membawa seluruh pemain di dalam kamar kembali ke lobi tanpa harus membuat ruangan baru dari awal.", bold_prefix="Fitur Rematch Instan: ")

    # SECTION 7
    add_h1("7. TANYA JAWAB (FAQ) & ETIKA BERMAIN")
    add_h2("FAQ 1: Mengapa Peramal hanya melihat 'WEREWOLF' atau 'NOT WEREWOLF'?")
    add_body(
        "Jawab: Ini adalah standar gameplay sosial deduksi kompetitif agar Peramal tidak mengetahui peran rekan timnya (seperti siapa Penjaga atau Pemburu) "
        "yang dapat merusak keseimbangan permainan deduksi."
    )
    
    add_h2("FAQ 2: Apa yang terjadi jika pemain menutup browser atau mengalami gangguan koneksi?")
    add_body(
        "Jawab: Sistem permainan dilengkapi mekanisme pemulihan sesi (Session Resume). Pemain yang menyegarkan (refresh) halaman "
        "akan langsung disambungkan kembali ke sesi kamar dan kartu peran pribadinya tanpa kehilangan progres."
    )

    add_h2("FAQ 3: Mengapa pengadilan tidak mengeksekusi siapa pun saat suara seri?")
    add_body(
        "Jawab: Hukum desa menerapkan asas praduga tidak bersalah. Eksekusi mati hanya dijatuhkan jika ada satu orang dengan suara kecurigaan mutlak terbanyak."
    )

    add_h2("Etika Bermain (Code of Conduct)")
    add_bullet("Dilarang melakukan screenshot peran pribadi untuk ditunjukkan di aplikasi perpesanan luar game (WhatsApp/Discord). Permainan berfokus pada seni berbicara dan berbohong.", bold_prefix="1. Kejujuran Deduksi : ")
    add_bullet("Tetaplah bersikap sopan dan santun selama sesi perdebatan. Jadikan perselisihan sebagai dinamika peran di dalam game.", bold_prefix="2. Sportivitas : ")
    add_bullet("Pemain yang sudah gugur dilarang membocorkan rahasia atau memberi petunjuk kepada pemain yang masih hidup (Ghosting).", bold_prefix="3. Hormati Korban Gugur : ")

    # Save Document
    output_filename = "c:/Users/Tude Raditya/werewolf-game/Panduan_Bermain_Werewolf.docx"
    doc.save(output_filename)
    print(f"Document successfully created at: {output_filename}")

if __name__ == "__main__":
    build_document()
